#!/usr/bin/env python3
"""Generate docs/WhatsApp_Assistant_Agent.pptx for LinkedIn / demos."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "WhatsApp_Assistant_Agent.pptx"

ACCENT = RGBColor(14, 165, 129)  # emerald
DARK = RGBColor(15, 23, 42)
MUTE = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)


def _box(slide, left, top, width, height, fill_rgb):
    shape = slide.shapes.add_shape(1, left, top, width, height)  # rectangle
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    shape.line.fill.background()
    return shape


def title_slide(prs, title: str, subtitle: str, footer: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    _box(slide, Inches(0), Inches(0), prs.slide_width, Inches(1.1), ACCENT)
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.5), Inches(9), Inches(1.2))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = DARK
    tb2 = slide.shapes.add_textbox(Inches(0.6), Inches(2.7), Inches(9), Inches(1))
    p2 = tb2.text_frame.paragraphs[0]
    p2.text = subtitle
    p2.font.size = Pt(22)
    p2.font.color.rgb = MUTE
    if footer:
        tb3 = slide.shapes.add_textbox(Inches(0.6), Inches(6.8), Inches(9), Inches(0.5))
        p3 = tb3.text_frame.paragraphs[0]
        p3.text = footer
        p3.font.size = Pt(14)
        p3.font.color.rgb = ACCENT


def bullet_slide(prs, title: str, bullets: list[str], note: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(9), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = DARK
    body = slide.shapes.add_textbox(Inches(0.75), Inches(1.35), Inches(8.8), Inches(5))
    tf = body.text_frame
    tf.word_wrap = True
    for i, line in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line
        para.level = 0
        para.font.size = Pt(20)
        para.font.color.rgb = DARK
        para.space_after = Pt(10)
    if note:
        n = slide.shapes.add_textbox(Inches(0.75), Inches(6.5), Inches(8.8), Inches(0.6))
        np = n.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(16)
        np.font.italic = True
        np.font.color.rgb = MUTE


def two_col_slide(
    prs,
    title: str,
    left_title: str,
    left_items: list[str],
    right_title: str,
    right_items: list[str],
    note: str = "",
):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(9), Inches(0.8))
    tb.text_frame.paragraphs[0].text = title
    tb.text_frame.paragraphs[0].font.size = Pt(32)
    tb.text_frame.paragraphs[0].font.bold = True
    tb.text_frame.paragraphs[0].font.color.rgb = DARK

    def col(x, heading, items):
        h = slide.shapes.add_textbox(x, Inches(1.3), Inches(4.2), Inches(0.5))
        hp = h.text_frame.paragraphs[0]
        hp.text = heading
        hp.font.size = Pt(18)
        hp.font.bold = True
        hp.font.color.rgb = ACCENT
        b = slide.shapes.add_textbox(x, Inches(1.85), Inches(4.2), Inches(4.5))
        tf = b.text_frame
        for i, line in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = line
            para.font.size = Pt(17)
            para.font.color.rgb = DARK
            para.space_after = Pt(8)

    col(Inches(0.75), left_title, left_items)
    col(Inches(5.1), right_title, right_items)
    if note:
        n = slide.shapes.add_textbox(Inches(0.75), Inches(6.5), Inches(8.8), Inches(0.6))
        np = n.text_frame.paragraphs[0]
        np.text = note
        np.font.size = Pt(16)
        np.font.italic = True
        np.font.color.rgb = MUTE


def table_slide(prs, title: str, rows: list[tuple[str, str]]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.45), Inches(9), Inches(0.8))
    tb.text_frame.paragraphs[0].text = title
    tb.text_frame.paragraphs[0].font.size = Pt(32)
    tb.text_frame.paragraphs[0].font.bold = True
    tb.text_frame.paragraphs[0].font.color.rgb = DARK

    nrows, ncols = len(rows) + 1, 2
    table = slide.shapes.add_table(nrows, ncols, Inches(0.6), Inches(1.35), Inches(8.8), Inches(0.45 * nrows)).table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(5.6)
    hdr = table.cell(0, 0), table.cell(0, 1)
    hdr[0].text = "Layer"
    hdr[1].text = "Technology"
    for cell in hdr:
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.size = Pt(14)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
        cell.text_frame.paragraphs[0].font.color.rgb = WHITE
    for r, (a, b) in enumerate(rows, start=1):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b
        for c in (0, 1):
            para = table.cell(r, c).text_frame.paragraphs[0]
            para.font.size = Pt(13)
            para.font.color.rgb = DARK


def closing_slide(prs, line1: str, line2: str, url: str = ""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _box(slide, Inches(0), Inches(6.2), prs.slide_width, Inches(1.3), ACCENT)
    t1 = slide.shapes.add_textbox(Inches(0.6), Inches(2.2), Inches(9), Inches(1))
    t1.text_frame.paragraphs[0].text = line1
    t1.text_frame.paragraphs[0].font.size = Pt(28)
    t1.text_frame.paragraphs[0].font.color.rgb = MUTE
    t2 = slide.shapes.add_textbox(Inches(0.6), Inches(3.2), Inches(9), Inches(1.2))
    p = t2.text_frame.paragraphs[0]
    p.text = line2
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = DARK
    if url:
        t3 = slide.shapes.add_textbox(Inches(0.6), Inches(6.45), Inches(9), Inches(0.5))
        t3.text_frame.paragraphs[0].text = url
        t3.text_frame.paragraphs[0].font.size = Pt(16)
        t3.text_frame.paragraphs[0].font.color.rgb = WHITE


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    title_slide(
        prs,
        "WhatsApp Customer Service Agent",
        "AI + human approval for logistics operations",
    )

    bullet_slide(
        prs,
        "The problem",
        [
            "Logistics CS runs on WhatsApp: ETAs, bookings, voice notes, parcel photos",
            "Slow first replies when handled manually",
            "Inconsistent tone across agents",
            "Risky auto-sends (promises, payments, private details)",
            "No single view for the operations head",
        ],
    )

    bullet_slide(
        prs,
        "The product",
        [
            "Company admin registers and adds customer-service agents",
            "Each agent links their business WhatsApp via QR scan",
            "AI drafts replies; sensitive items go to a hold queue",
            "Head sees all messages, held items, and activity",
            "Text, voice notes, and images use one pipeline",
        ],
    )

    two_col_slide(
        prs,
        "How a reply is decided",
        "Auto-send when",
        [
            "Reply is not empty",
            "Facts are grounded (KB / schedule)",
            "Not a human-only decision",
        ],
        "Hold when",
        [
            "Meetings / scheduling commitments",
            "Payments or private details",
            "Model is unsure (fail closed)",
        ],
        note="Doubt defaults to hold — customers never get a guessed promise.",
    )

    bullet_slide(
        prs,
        "Dashboard",
        [
            "Chat — internal assistant (Mika): draft, send, read messages",
            "Held — approve, edit, or dismiss with reason",
            "Messages — full in/out log by contact and day",
            "Activity — audit trail",
            "Team — admin manages agents",
            "WhatsApp — QR linking per agent",
        ],
        note="Global toggles: auto-reply and direct send.",
    )

    bullet_slide(
        prs,
        "Architecture",
        [
            "Customer → WhatsApp Web (one session per agent)",
            "Node bridge (whatsapp-web.js) → FastAPI /webhook",
            "Voice → faster-whisper · Photos → BLIP (Transformers + PyTorch)",
            "LangGraph agent → OpenAI (reply + safety flags)",
            "Send on WhatsApp OR store in SQLite held queue",
            "Operations team uses the web dashboard",
        ],
    )

    table_slide(
        prs,
        "Technology stack",
        [
            ("API & dashboard", "FastAPI, Uvicorn, Pydantic"),
            ("Orchestration", "LangGraph + LangChain"),
            ("LLM / TTS", "OpenAI"),
            ("Data", "SQLite"),
            ("Speech-to-text", "faster-whisper"),
            ("Vision", "Hugging Face Transformers, PyTorch, BLIP"),
            ("WhatsApp", "whatsapp-web.js (multi-agent QR)"),
            ("Auth", "Admin + agent accounts, scoped history"),
        ],
    )

    bullet_slide(
        prs,
        "Why logistics customer service",
        [
            "First response in seconds for routine questions",
            "Humans in the loop for bookings and commitments",
            "One ops view for the company head",
            "Each agent owns a WhatsApp line",
            "English & Urdu (including Roman Urdu); voice in → voice out",
        ],
    )

    closing_slide(
        prs,
        'Not "reply to everything."',
        "Reply like the company would — and hold the rest.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
